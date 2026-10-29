"""Case Alerts web app: invite-only login, upload Caseboard PDF, download workbook."""
from __future__ import annotations

import json
import os
import secrets
import shutil
import threading
import time
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import openpyxl
from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, RedirectResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from starlette.middleware.sessions import SessionMiddleware

from . import db
from .bulk import fetch as bulk_fetch
from .bulk import job as bulk
from .pipeline import run as run_pipeline
from .workbook import HEADERS

APP_NAME = os.environ.get("APP_NAME", "Case Alerts")
BASE_URL = os.environ.get("BASE_URL", "").rstrip("/")
MAX_PDF = 25 * 1024 * 1024
TEMPLATE_PATH = db.DATA / "template.xlsx"

app = FastAPI(docs_url=None, redoc_url=None, openapi_url=None)
app.add_middleware(
    SessionMiddleware,
    secret_key=os.environ.get("SECRET_KEY") or secrets.token_hex(32),
    session_cookie="ca_session", max_age=60 * 60 * 12, same_site="lax",
    https_only=os.environ.get("COOKIE_SECURE", "1") == "1",
)
app.mount("/static", StaticFiles(directory=str(Path(__file__).parent / "static")), name="static")
templates = Jinja2Templates(directory=str(Path(__file__).parent / "templates"))
_LONDON = ZoneInfo("Europe/London")


def _nice_date(v, with_time=False):
    """'2026-10-16' -> '16 Oct 2026'; ISO timestamps are shown in UK time."""
    if not v:
        return ""
    try:
        d = datetime.fromisoformat(str(v))
    except ValueError:
        return v
    if d.tzinfo:
        d = d.astimezone(_LONDON)
    out = f"{d.day} {d:%b %Y}"
    return out + (f" at {d:%H:%M}" if with_time else "")


templates.env.filters["nice_date"] = _nice_date
pool = ThreadPoolExecutor(max_workers=2)
progress: dict[str, dict] = {}
_fails: dict[str, list[float]] = defaultdict(list)


@app.on_event("startup")
def startup():
    db.init()
    bulk.init()
    email = os.environ.get("ADMIN_EMAIL")
    if os.environ.get("ADMIN_PASSWORD"):
        db.ensure_admin(email, os.environ["ADMIN_PASSWORD"], os.environ.get("ADMIN_NAME", "Administrator"))
    elif email:
        # No password in config: create the admin without one and issue a single-use setup link.
        u = db.user_by_email(email)
        if not u:
            with db.conn() as c:
                c.execute("INSERT INTO users(email,name,pw_hash,is_admin,created) VALUES(?,?,NULL,1,?)",
                          (email.strip().lower(), os.environ.get("ADMIN_NAME", "Administrator"), db.now()))
            u = db.user_by_email(email)
        if not u["pw_hash"]:
            saved = db.DATA / "setup-token.txt"
            token = saved.read_text().strip() if saved.exists() else ""
            if not token or not db.get_invite(token):
                with db.conn() as c:
                    c.execute("UPDATE invites SET used='superseded' WHERE email=? AND used IS NULL", (u["email"],))
                token = db.make_invite(u["email"], u["id"], kind="setup", days=2)
                saved.write_text(token)
            path = f"/invite/{token}"
            (db.DATA / "setup-link.txt").write_text((BASE_URL or "") + path + "\n")
            print(f"ADMIN SETUP LINK: {(BASE_URL or '')}{path}", flush=True)
    # Invites requested by the owner through the INVITE_EMAILS setting (comma separated).
    # Each gets one stable link, logged at start-up, until it is used or expires.
    admin = db.user_by_email(email) if email else None
    tokens_dir = db.DATA / "invite-tokens"
    tokens_dir.mkdir(exist_ok=True)
    for addr in [a.strip().lower() for a in os.environ.get("INVITE_EMAILS", "").split(",") if a.strip()]:
        if db.user_by_email(addr):
            continue
        f = tokens_dir / addr
        token = f.read_text().strip() if f.exists() else ""
        if not token or not db.get_invite(token):
            token = db.make_invite(addr, admin["id"] if admin else 0, kind="invite", days=7)
            f.write_text(token)
        print(f"INVITE LINK {addr}: {(BASE_URL or '')}/invite/{token}", flush=True)
    # Password resets requested by the owner through RESET_EMAILS (comma separated).
    # One stable single-use link per address, logged at start-up; no password is ever set here.
    reset_dir = db.DATA / "reset-tokens"
    reset_dir.mkdir(exist_ok=True)
    for addr in [a.strip().lower() for a in os.environ.get("RESET_EMAILS", "").split(",") if a.strip()]:
        u = db.user_by_email(addr)
        if not u:
            print(f"RESET: no account for {addr}", flush=True)
            continue
        f = reset_dir / addr
        token = f.read_text().strip() if f.exists() else ""
        if not token or not db.get_invite(token):
            token = db.make_invite(addr, u["id"], kind="reset", days=2)
            f.write_text(token)
        print(f"RESET LINK {addr}: {(BASE_URL or '')}/invite/{token}", flush=True)
    # runs interrupted by a restart
    for r in db.list_runs(limit=1000):
        if r["status"] in ("queued", "running"):
            db.update_run(r["id"], status="failed", error="Interrupted by a server restart. Please upload again.")


# ---- helpers --------------------------------------------------------------------
def current_user(request: Request):
    uid = request.session.get("uid")
    if not uid:
        return None
    u = db.user_by_id(uid)
    return u if u and u["active"] else None


def need_user(request: Request):
    u = current_user(request)
    if not u:
        raise HTTPException(status_code=303, headers={"Location": "/login"})
    return u


def need_admin(request: Request):
    u = need_user(request)
    if not u["is_admin"]:
        raise HTTPException(status_code=403)
    return u


def csrf(request: Request) -> str:
    if "csrf" not in request.session:
        request.session["csrf"] = secrets.token_urlsafe(24)
    return request.session["csrf"]


def check_csrf(request: Request, token: str):
    if not token or token != request.session.get("csrf"):
        raise HTTPException(status_code=400, detail="Form expired. Go back and try again.")


def page(request: Request, _tpl: str, **ctx):
    ctx.update(request=request, user=current_user(request), csrf=csrf(request), app_name=APP_NAME,
               template_ready=TEMPLATE_PATH.exists(), ch_ready=bool(os.environ.get("CH_API_KEY")))
    return templates.TemplateResponse(request, _tpl, ctx)


def link(request: Request, path: str) -> str:
    return (BASE_URL or str(request.base_url).rstrip("/")) + path


def throttled(key: str) -> bool:
    now = time.time()
    _fails[key] = [t for t in _fails[key] if now - t < 900]
    return len(_fails[key]) >= 8


def check_template(path: Path) -> str | None:
    try:
        wb = openpyxl.load_workbook(path, read_only=True)
    except Exception:
        return "That file isn't a readable .xlsx workbook."
    if "Sheet1" not in wb.sheetnames:
        return "The template needs a sheet named Sheet1."
    ws = wb["Sheet1"]
    hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1, max_col=14))]
    if [str(h).strip() if h else h for h in hdr] != HEADERS:
        return "Row 1 of Sheet1 must be exactly: " + " | ".join(HEADERS)
    return None


# ---- auth -------------------------------------------------------------------------
@app.get("/login", response_class=HTMLResponse)
def login_form(request: Request):
    if current_user(request):
        return RedirectResponse("/", 303)
    return page(request, "login.html")


@app.post("/login")
def login(request: Request, email: str = Form(...), password: str = Form(...), csrf_token: str = Form("")):
    check_csrf(request, csrf_token)
    key = f"{request.client.host if request.client else ''}|{email.lower()}"
    if throttled(key):
        return page(request, "login.html", error="Too many attempts. Try again in 15 minutes.", email=email)
    u = db.user_by_email(email)
    if not u or not u["active"] or not db.check_pw(password, u["pw_hash"]):
        _fails[key].append(time.time())
        return page(request, "login.html", error="Email or password not recognised.", email=email)
    request.session.clear()
    request.session["uid"] = u["id"]
    db.touch_login(u["id"])
    return RedirectResponse("/", 303)


@app.post("/logout")
def logout(request: Request, csrf_token: str = Form("")):
    check_csrf(request, csrf_token)
    request.session.clear()
    return RedirectResponse("/login", 303)


@app.get("/invite/{token}", response_class=HTMLResponse)
def invite_form(request: Request, token: str):
    inv = db.get_invite(token)
    if not inv:
        return page(request, "message.html", title="Link expired",
                    message="This link is invalid, used or expired. Ask your administrator for a new one.")
    existing = db.user_by_email(inv["email"])
    return page(request, "invite.html", inv=inv, token=token, reset=bool(existing))


@app.post("/invite/{token}")
def invite_accept(request: Request, token: str, name: str = Form(""), password: str = Form(...),
                  password2: str = Form(...), csrf_token: str = Form("")):
    check_csrf(request, csrf_token)
    inv = db.get_invite(token)
    if not inv:
        return RedirectResponse(f"/invite/{token}", 303)
    existing = db.user_by_email(inv["email"])
    err = None
    if len(password) < 10:
        err = "Use at least 10 characters."
    elif password != password2:
        err = "The passwords don't match."
    elif not existing and not name.strip():
        err = "Enter your name."
    if err:
        return page(request, "invite.html", inv=inv, token=token, reset=bool(existing), error=err, name=name)
    if existing:
        db.set_password(existing["id"], password)
        db.set_flag(existing["id"], "active", True)
        uid = existing["id"]
    else:
        db.create_user(inv["email"], name, password)
        uid = db.user_by_email(inv["email"])["id"]
    db.use_invite(token)
    request.session.clear()
    request.session["uid"] = uid
    return RedirectResponse("/", 303)


# ---- main tool ------------------------------------------------------------------
@app.get("/", response_class=HTMLResponse)
def home(request: Request):
    u = need_user(request)
    runs = db.list_runs(None if u["is_admin"] else u["id"], limit=30)
    return page(request, "home.html", runs=[_run_view(r) for r in runs])


def _run_view(r):
    s = json.loads(r["summary"]) if r["summary"] else None
    return {**dict(r), "s": s}


@app.post("/runs")
async def create_run(request: Request, pdf: UploadFile = File(...), template: UploadFile | None = File(None),
                     csrf_token: str = Form("")):
    u = need_user(request)
    check_csrf(request, csrf_token)
    if not os.environ.get("CH_API_KEY"):
        return page(request, "message.html", title="Not ready", message="The Companies House API key is not configured.")
    data = await pdf.read()
    if len(data) > MAX_PDF or not data.startswith(b"%PDF"):
        return page(request, "message.html", title="Wrong file",
                    message="Upload the Caseboard alert as a PDF (max 25 MB).")
    rid = db.new_run(u["id"], pdf.filename or "alert.pdf")
    rdir = db.DATA / "runs" / rid
    rdir.mkdir(parents=True)
    (rdir / "input.pdf").write_bytes(data)

    tmpl = TEMPLATE_PATH
    if template is not None and template.filename:
        tdata = await template.read()
        tpath = rdir / "template.xlsx"
        tpath.write_bytes(tdata)
        err = check_template(tpath)
        if err:
            db.update_run(rid, status="failed", error=f"Template rejected: {err}")
            return RedirectResponse(f"/runs/{rid}", 303)
        tmpl = tpath
    if not tmpl.exists():
        db.update_run(rid, status="failed", error="No template uploaded yet. An admin needs to add it under Admin.")
        return RedirectResponse(f"/runs/{rid}", 303)

    progress[rid] = {"done": 0, "total": 0, "current": "Reading PDF"}
    pool.submit(_work, rid, rdir, tmpl)
    return RedirectResponse(f"/runs/{rid}", 303)


_busy = {"n": 0}
_busy_lock = threading.Lock()


def _alerts_busy(delta):
    """While an alert is processing, the bulk sheet fill slows right down so the alert gets the API."""
    with _busy_lock:
        _busy["n"] = max(0, _busy["n"] + delta)
        (bulk_fetch.ALERTS_BUSY.set if _busy["n"] else bulk_fetch.ALERTS_BUSY.clear)()


def _work(rid: str, rdir: Path, tmpl: Path):
    _alerts_busy(+1)
    db.update_run(rid, status="running")

    def prog(done, total, case):
        progress[rid] = {"done": done, "total": total, "current": case.subject_name}

    try:
        tcopy = rdir / "_template_copy.xlsx"
        shutil.copyfile(tmpl, tcopy)  # never touch the master template
        s = run_pipeline(rdir / "input.pdf", tcopy, rdir / "out", history=db.history_for, progress=prog)
        tcopy.unlink(missing_ok=True)
        db.record_filings(rid, s["alert_date"], s.pop("company_numbers"))
        s["xlsx"] = Path(s["xlsx"]).name
        s["register"] = Path(s["register"]).name
        db.update_run(rid, status="done", alert_date=s["alert_date"], summary=s)
    except Exception as e:
        db.update_run(rid, status="failed", error=f"{type(e).__name__}: {e}")
    finally:
        progress.pop(rid, None)
        _alerts_busy(-1)


def _own_run(request: Request, rid: str):
    u = need_user(request)
    r = db.get_run(rid)
    if not r or (r["user_id"] != u["id"] and not u["is_admin"]):
        raise HTTPException(status_code=404)
    return r


@app.get("/runs/{rid}", response_class=HTMLResponse)
def run_page(request: Request, rid: str):
    r = _own_run(request, rid)
    return page(request, "run.html", r=_run_view(r), p=progress.get(rid))


@app.get("/runs/{rid}/status")
def run_status(request: Request, rid: str):
    r = _own_run(request, rid)
    return JSONResponse({"status": r["status"], "progress": progress.get(rid)})


@app.get("/runs/{rid}/download/{kind}")
def download(request: Request, rid: str, kind: str):
    r = _own_run(request, rid)
    if r["status"] != "done":
        raise HTTPException(404)
    s = json.loads(r["summary"])
    name = {"xlsx": s["xlsx"], "register": s["register"]}.get(kind)
    if not name:
        raise HTTPException(404)
    path = db.DATA / "runs" / rid / "out" / name
    media = ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" if kind == "xlsx" else "text/csv")
    return FileResponse(path, media_type=media, filename=name)


# ---- admin ------------------------------------------------------------------------
@app.get("/admin", response_class=HTMLResponse)
def admin(request: Request):
    need_admin(request)
    return page(request, "admin.html", users=db.all_users(), invites=db.open_invites(),
                new_link=request.session.pop("new_link", None), notice=request.session.pop("notice", None))


@app.post("/admin/invite")
def admin_invite(request: Request, email: str = Form(...), csrf_token: str = Form("")):
    u = need_admin(request)
    check_csrf(request, csrf_token)
    email = email.strip().lower()
    existing = db.user_by_email(email)
    token = db.make_invite(email, u["id"], kind="reset" if existing else "invite",
                           days=2 if existing else 7)
    request.session["new_link"] = {"email": email, "url": link(request, f"/invite/{token}"),
                                   "kind": "password reset" if existing else "invite"}
    return RedirectResponse("/admin", 303)


@app.post("/admin/invites/{rowid}/revoke")
def admin_revoke(request: Request, rowid: int, csrf_token: str = Form("")):
    need_admin(request)
    check_csrf(request, csrf_token)
    db.revoke_invite(rowid)
    return RedirectResponse("/admin", 303)


@app.post("/admin/users/{uid}/{action}")
def admin_user(request: Request, uid: int, action: str, csrf_token: str = Form("")):
    me = need_admin(request)
    check_csrf(request, csrf_token)
    if uid == me["id"]:
        request.session["notice"] = "You can't change your own account here."
        return RedirectResponse("/admin", 303)
    if action in ("disable", "enable"):
        db.set_flag(uid, "active", action == "enable")
    elif action in ("make-admin", "remove-admin"):
        db.set_flag(uid, "is_admin", action == "make-admin")
    return RedirectResponse("/admin", 303)


@app.post("/admin/template")
async def admin_template(request: Request, template: UploadFile = File(...), csrf_token: str = Form("")):
    need_admin(request)
    check_csrf(request, csrf_token)
    tmp = db.DATA / "template.upload.xlsx"
    tmp.write_bytes(await template.read())
    err = check_template(tmp)
    if err:
        tmp.unlink(missing_ok=True)
        request.session["notice"] = f"Template rejected: {err}"
    else:
        tmp.replace(TEMPLATE_PATH)
        request.session["notice"] = "Template saved. New runs will use it."
    return RedirectResponse("/admin", 303)


@app.get("/admin/template")
def admin_template_download(request: Request):
    need_admin(request)
    if not TEMPLATE_PATH.exists():
        raise HTTPException(404)
    return FileResponse(TEMPLATE_PATH, filename="Case Alerts - template.xlsx")


# ---- insolvency sheet fill (admin) ---------------------------------------------
@app.get("/admin/bulk", response_class=HTMLResponse)
def bulk_page(request: Request):
    need_admin(request)
    jobs = []
    for j in bulk.all_jobs():
        d = dict(j)
        d["stats"] = json.loads(j["stats"]) if j["stats"] else None
        left = (d["total"] - d["done"]) / d["rate"] / 3600 if d.get("rate") and d["status"] == "running" else None
        d["hours_left"] = round(left, 1) if left is not None else None
        jobs.append(d)
    return page(request, "bulk.html", jobs=jobs, notice=request.session.pop("notice", None))


@app.post("/admin/bulk")
async def bulk_create(request: Request, workbook: UploadFile = File(...), progress_file: UploadFile | None = File(None),
                      csrf_token: str = Form("")):
    u = need_admin(request)
    check_csrf(request, csrf_token)
    data = await workbook.read()
    if not data[:2] == b"PK" or len(data) > 40 * 1024 * 1024:
        request.session["notice"] = "Upload the sheet as an .xlsx file (max 40 MB)."
        return RedirectResponse("/admin/bulk", 303)
    prog = None
    if progress_file is not None and progress_file.filename:
        prog = await progress_file.read()
    try:
        bulk.create(u["id"], workbook.filename or "sheet.xlsx", data, prog)
        request.session["notice"] = "Started. You can leave this page; it carries on in the background."
    except ValueError as e:
        request.session["notice"] = str(e)
    return RedirectResponse("/admin/bulk", 303)


@app.get("/admin/bulk/{jid}/status")
def bulk_status(request: Request, jid: str):
    need_admin(request)
    j = bulk.get(jid)
    if not j:
        raise HTTPException(404)
    return JSONResponse({k: j[k] for k in ("status", "total", "done", "rate", "error")})


@app.post("/admin/bulk/{jid}/{action}")
def bulk_action(request: Request, jid: str, action: str, csrf_token: str = Form("")):
    need_admin(request)
    check_csrf(request, csrf_token)
    if not bulk.get(jid):
        raise HTTPException(404)
    if action == "stop":
        bulk.stop(jid)
    elif action == "resume":
        bulk.resume(jid)
    return RedirectResponse("/admin/bulk", 303)


@app.get("/admin/bulk/{jid}/download/{kind}")
def bulk_download(request: Request, jid: str, kind: str, partial: int = 0):
    need_admin(request)
    j = bulk.get(jid)
    if not j or kind not in ("xlsx", "sources"):
        raise HTTPException(404)
    stem = Path(j["filename"] or "sheet.xlsx").stem
    if j["status"] == "done" and not partial:
        d = bulk.folder(jid)
        out, log_csv = d / "filled.xlsx", d / "filled-sources.csv"
        if not out.exists():
            out, log_csv, _ = bulk.build_output(jid)
        label = "filled"
    else:
        out, log_csv, _ = bulk.build_output(jid, partial=True)
        label = f"partial {j['done']} of {j['total']}"
    if kind == "xlsx":
        return FileResponse(out, filename=f"{stem} ({label}).xlsx",
                            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    return FileResponse(log_csv, filename=f"{stem} sources ({label}).csv", media_type="text/csv")


@app.get("/healthz")
def healthz():
    return {"ok": True}


@app.exception_handler(HTTPException)
async def http_errors(request: Request, exc: HTTPException):
    if exc.status_code == 303:
        return RedirectResponse(exc.headers["Location"], 303)
    msg = {403: "You don't have access to that page.", 404: "Not found."}.get(exc.status_code, exc.detail)
    return page(request, "message.html", title="Sorry", message=msg)
