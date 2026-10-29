"""Fill an insolvency workbook from the job's findings. Only blank cells are written, so
anything typed by hand stays. Columns are found by their header, on any sheet."""
from __future__ import annotations

import csv
import json
import re
from collections import Counter
from datetime import date

import openpyxl

from .resolve import resolve

NUMBER_HEADERS = {"company number", "company_number", "company no", "company no.", "crn", "registered number"}


def norm(v) -> str:
    v = re.sub(r"\s+", "", str(v)).upper()
    return v.zfill(8) if v.isdigit() else v


def uk(d):
    try:
        return date.fromisoformat(d).strftime("%d/%m/%Y")
    except (TypeError, ValueError):
        return ""


def _hkey(h):
    return re.sub(r"[^a-z0-9]+", " ", str(h or "").lower()).strip()


def number_column(ws):
    for i, c in enumerate(next(ws.iter_rows(min_row=1, max_row=1)), start=1):
        if _hkey(c.value) in {_hkey(x) for x in NUMBER_HEADERS}:
            return i
    return None


def company_numbers(path) -> list[str]:
    """Every company number in the workbook, in order, without repeats."""
    wb = openpyxl.load_workbook(path, read_only=True)
    seen = {}
    for ws in wb.worksheets:
        col = number_column(ws)
        if not col:
            continue
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col, values_only=True):
            v = row[0]
            if v not in (None, ""):
                n = norm(v)
                if re.fullmatch(r"[A-Z]{0,2}\d{5,8}", n):
                    seen.setdefault(n, None)
    return list(seen)


def load_findings(cache_path) -> dict:
    found = {}
    with open(cache_path) as fh:
        for line in fh:
            try:
                rec = json.loads(line)
            except ValueError:
                continue
            found[rec["n"]] = rec
    return found


def values_for(o) -> dict:
    """Header (normalised) -> value, covering both layouts in the CCO sheet."""
    ips = o["ips"]
    return {
        "insolvency practitioner 1": ips[0] if len(ips) > 0 else None,
        "insolvency practitioner 2": ips[1] if len(ips) > 1 else None,
        "insolvency practitioner 3": ips[2] if len(ips) > 2 else None,
        "insolvency practitioners": "; ".join(ips) or None,
        "insolvency firm": o["firm"] or None,
        "insolvency practitioners firm": o["firm"] or None,
        "unsatisfied charge holder": o["unsatisfied"] or None,
        "unsatisfied charge holders": o["unsatisfied"] or None,
        "date registered": ("; ".join(uk(d) or "no date" for d in o["unsat_dates"].split("; "))
                            if o["unsatisfied"] else None),
        "satisfied charge": o["satisfied"] or None,
        "date satisfied": ("; ".join(uk(d) or "no date" for d in o["sat_dates"].split("; "))
                           if o["satisfied"] else None),
    }


def fill(src, cache_path, out_path, log_path) -> dict:
    found = load_findings(cache_path)
    res = {n: resolve(r) for n, r in found.items()}
    wb = openpyxl.load_workbook(src)
    stats = Counter()
    for ws in wb.worksheets:
        ncol = number_column(ws)
        if not ncol:
            continue
        headers = {i: _hkey(c.value) for i, c in enumerate(ws[1], start=1)}
        targets = {i: h for i, h in headers.items() if h in values_for({"ips": [], "firm": "", "unsatisfied": "",
                                                                         "unsat_dates": "", "satisfied": "",
                                                                         "sat_dates": ""})}
        if not targets:
            continue
        for r in range(2, ws.max_row + 1):
            v = ws.cell(r, ncol).value
            if v in (None, ""):
                continue
            o = res.get(norm(v))
            if o is None:
                stats["rows not fetched yet"] += 1
                continue
            vals = values_for(o)
            for i, h in targets.items():
                c = ws.cell(r, i)
                if c.value in (None, "") and vals[h]:
                    c.value = vals[h]
            stats["rows filled"] += 1
    wb.save(out_path)

    with open(log_path, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["company_number", "practitioners", "firm", "unsatisfied_charge_holders", "sources",
                    "companies_house_insolvency", "companies_house_charges", "gazette", "gazette_notice"])
        for n, o in res.items():
            gz = (found[n].get("gazette") or {}).get("notices") or []
            w.writerow([n, "; ".join(o["ips"]), o["firm"], o["unsatisfied"], "; ".join(o["sources"]),
                        o["status"]["insolvency"], o["status"]["charges"], o["status"]["gazette"],
                        f"https://www.thegazette.co.uk/notice/{gz[0]['id']}" if gz else ""])
            stats["companies"] += 1
            stats["with practitioners"] += bool(o["ips"])
            stats["with firm"] += bool(o["firm"])
            stats["with charge holders"] += bool(o["unsatisfied"])
    return dict(stats)
