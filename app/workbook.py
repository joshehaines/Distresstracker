"""Template-preserving workbook writer + programmatic validation."""
from __future__ import annotations

import csv
import shutil
from copy import copy
from datetime import datetime
from pathlib import Path

import openpyxl

from .companies_house import Lookup, NA, VERIFIED, norm

HEADERS = [
    "Date", "Company Name", "Filing Type", "Petitioner", "SIC Codes", "SIC Description(s)",
    "Multiple Caseboard Filings", "Property Postcode", "Trading Name", "Charge Holder",
    "Property ID (IBIS)", "Property name", "Notes/Outcome", "Action",
]


def build_rows(cases, lookups: dict[str, Lookup], history: dict[str, list[str]] | None = None):
    """history: company_number -> earlier case refs (from previous runs)."""
    history = history or {}
    by_number: dict[str, list[str]] = {}
    for c in cases:
        lk = lookups[c.case_ref]
        if lk.company_number:
            by_number.setdefault(lk.company_number, []).append(c.case_ref)

    rows = []
    for c in cases:
        lk = lookups[c.case_ref]
        notes: list[str] = []
        if c.subject_type == "individual":
            notes.append("Not applicable: individual proceeding.")
        elif lk.status == VERIFIED:
            s = f"SIC Verified: CH {lk.company_number}"
            if lk.company_status:
                s += f" ({lk.company_status.replace('-', ' ')})"
            notes.append(s + f", checked {lk.checked}.")
            if lk.legal_name and norm(lk.legal_name) != norm(c.search_name):
                notes.append(f"Registered name: {lk.legal_name}.")
            if lk.reason:
                notes.append(lk.reason)
        else:
            notes.append(f"SIC {lk.status}: {lk.reason.rstrip('. ')}.")
        notes.extend(c.notes)

        multiple = ""
        if lk.company_number:
            same_run = [r for r in by_number.get(lk.company_number, []) if r != c.case_ref]
            earlier = [r for r in history.get(lk.company_number, []) if r != c.case_ref]
            others = same_run + earlier
            if others:
                multiple = "Yes"
                notes.append("Other filings: " + ", ".join(dict.fromkeys(others)) + ".")

        charge = ""
        if lk.charge_status == VERIFIED:
            if lk.charge_holders:
                charge = "; ".join(lk.charge_holders)
            elif lk.charges_total == 0:
                charge = "No charges"
            else:
                charge = "No outstanding charges"
        elif lk.status == VERIFIED and lk.charge_status:
            notes.append("Charge register could not be retrieved.")

        rows.append({
            "case_ref": c.case_ref,
            "subject_type": c.subject_type,
            "status": lk.status,
            "values": [
                datetime.fromisoformat(c.alert_date) if c.alert_date else None,
                c.subject_name,
                c.filing_type,
                c.petitioner or None,
                ";".join(lk.sic_codes) or None,
                ";".join(lk.sic_descriptions) or None,
                multiple or None,
                None,               # Property Postcode: not auto-verifiable
                None,               # Trading Name
                charge or None,
                None,               # Property ID (IBIS)
                None,               # Property name
                " ".join(n.strip() for n in notes if n.strip()) or None,
                None,               # Action: left for the user
            ],
        })
    return rows


def write_workbook(template: str | Path, out: str | Path, rows) -> None:
    template, out = Path(template), Path(out)
    if template.resolve() == out.resolve():
        raise ValueError("Refusing to overwrite the template")
    shutil.copyfile(template, out)
    wb = openpyxl.load_workbook(out)
    ws = wb["Sheet1"]
    for extra in [n for n in wb.sheetnames if n != "Sheet1"]:
        del wb[extra]

    # data-row style taken from the template's first example row
    style_row = 2
    styles = []
    for col in range(1, 15):
        src = ws.cell(style_row, col)
        styles.append(dict(font=copy(src.font), fill=copy(src.fill), border=copy(src.border),
                           alignment=copy(src.alignment), protection=copy(src.protection),
                           number_format=src.number_format))
    styles[0]["number_format"] = "d-mmm-yy"
    row_height = ws.row_dimensions[style_row].height

    # header exactly as specified
    for i, h in enumerate(HEADERS, start=1):
        ws.cell(1, i).value = h

    # remove all historical rows and their row formatting
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    for r in list(ws.row_dimensions.keys()):
        if r > 1:
            del ws.row_dimensions[r]

    for r, row in enumerate(rows, start=2):
        for col, val in enumerate(row["values"], start=1):
            cell = ws.cell(r, col)
            st = styles[col - 1]
            cell.font, cell.fill, cell.border = copy(st["font"]), copy(st["fill"]), copy(st["border"])
            cell.alignment, cell.protection = copy(st["alignment"]), copy(st["protection"])
            cell.number_format = st["number_format"]
            if col == 5 and val is not None:
                cell.value = str(val)
                cell.number_format = "@"
            else:
                cell.value = val
        if row_height:
            ws.row_dimensions[r].height = row_height

    ws.freeze_panes = "A2"
    wb.save(out)


def validate(out: str | Path, template: str | Path, template_mtime: float, cases, rows) -> list[str]:
    """Reopen the saved workbook and check it against the spec. Returns failures."""
    fails: list[str] = []
    wb = openpyxl.load_workbook(out)
    if wb.sheetnames != ["Sheet1"]:
        fails.append(f"Sheets are {wb.sheetnames}, expected only Sheet1")
    ws = wb["Sheet1"]
    hdr = [ws.cell(1, i).value for i in range(1, 15)]
    if hdr != HEADERS:
        fails.append(f"Header mismatch: {hdr}")
    if ws.max_column != 14:
        fails.append(f"Used columns = {ws.max_column}, expected 14")
    expected_last = len(rows) + 1
    if ws.max_row != expected_last:
        fails.append(f"Used range ends at row {ws.max_row}, expected {expected_last}")

    refs = [c.case_ref for c in cases]
    if len(rows) != len(set(refs)):
        fails.append(f"{len(rows)} output rows vs {len(set(refs))} distinct case refs")

    for i, (c, row) in enumerate(zip(cases, rows), start=2):
        vals = [ws.cell(i, k).value for k in range(1, 15)]
        if not isinstance(vals[0], datetime):
            fails.append(f"Row {i}: date is not a real date")
        if vals[1] != c.subject_name:
            fails.append(f"Row {i}: company name mismatch")
        e, f = vals[4], vals[5]
        if e is not None and not isinstance(e, str):
            fails.append(f"Row {i}: SIC codes not stored as text")
        if (e is None) != (f is None) or (e and len(e.split(";")) != len(f.split(";"))):
            fails.append(f"Row {i}: SIC codes/descriptions misaligned")
        if c.subject_type == "individual":
            if e or f or "Not applicable: individual proceeding" not in (vals[12] or ""):
                fails.append(f"Row {i}: individual not marked SIC not applicable")
        elif c.subject_type == "company":
            if not e and not (vals[12] or "").startswith("SIC "):
                fails.append(f"Row {i}: company has no SIC data and no failure reason")
        else:
            if not e and not (vals[12] or "").strip():
                fails.append(f"Row {i}: unidentified subject without explanation")
        for k in (7, 8, 10, 11, 13):
            if vals[k] is not None:
                fails.append(f"Row {i}: unverifiable column {HEADERS[k]} populated")

    if Path(template).stat().st_mtime != template_mtime:
        fails.append("Template file was modified")
    return fails


def write_register(path: str | Path, cases, lookups) -> None:
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["order", "page", "alert_date", "case_ref", "filed", "proceeding", "subject_source",
                    "search_name", "subject_type", "supplied_company_number", "petitioner",
                    "lookup_status", "verified_company_number", "registered_name", "company_status",
                    "profile_url", "checked", "sic_codes", "charge_status", "charge_holders",
                    "attempts", "reason", "notes"])
        for c in cases:
            lk = lookups[c.case_ref]
            w.writerow([c.order, c.page, c.alert_date, c.case_ref, c.filed, c.proceeding, c.heading,
                        c.search_name, c.subject_type, c.company_number or "", c.petitioner,
                        lk.status, lk.company_number or "", lk.legal_name or "", lk.company_status or "",
                        lk.profile_url or "", lk.checked or "", ";".join(lk.sic_codes),
                        lk.charge_status or "", "; ".join(lk.charge_holders),
                        " | ".join(lk.attempts), lk.reason, " ".join(c.notes)])
